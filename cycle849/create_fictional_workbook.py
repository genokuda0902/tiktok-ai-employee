"""Generate a tiny fictional workbook for offline LibreOffice Calc capture.

No personal, employee, customer or real sales data is used.
Requires openpyxl (free). Capture the opened workbook locally with X11/FFmpeg;
never commit or upload the screen recording before human rights review.
"""
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from pathlib import Path
import argparse

def create(path):
    w=Workbook();s=w.active;s.title='架空データ_集計例'
    s.sheet_view.zoomScale=140;s.sheet_view.showGridLines=False
    rows=[['月','営業A','営業B'],['4月',12,8],['5月',16,9],['6月',18,12],
          ['合計','=SUM(B2:B4)','=SUM(C2:C4)']]
    for row in rows:s.append(row)
    for col in ('A','B','C'):s.column_dimensions[col].width=18
    for r in range(1,6):
        s.row_dimensions[r].height=43
        for c in range(1,4):
            cell=s.cell(r,c)
            cell.alignment=Alignment(horizontal='center',vertical='center')
            cell.font=Font(name='Noto Sans CJK JP',size=20,bold=r in (1,5),
                           color='FFFFFF' if r in (1,5) else '172C40')
            color='073F52' if r==1 else ('076F70' if r==5 else ('EAF4F5' if r%2==0 else 'FFFFFF'))
            cell.fill=PatternFill('solid',fgColor=color)
            cell.border=Border(bottom=Side(style='thin',color='C5D7DC'))
    s.freeze_panes='A2'
    s['A7']='架空の例：実測の売上ではありません'
    s['A7'].font=Font(name='Noto Sans CJK JP',size=13,color='4B6473')
    s.sheet_view.selection[0].activeCell='B5'
    s.sheet_view.selection[0].sqref='B5'
    path=Path(path)
    if path.exists():raise FileExistsError(path)
    path.parent.mkdir(parents=True,exist_ok=True)
    w.save(path)
    return path

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output',required=True)
    print(create(ap.parse_args().output))
